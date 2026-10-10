"""
Auto Shutdown Module
Supports: Shutdown, Restart, Hibernate, Logoff, Standby, Lock,
          Close App, Run App, Message
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import subprocess
import os
import sys
import threading
import time
import json
import re
import csv
import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from constants import COLORS, FONTS


class AutoShutdown:
    def __init__(self, parent):
        self.parent = parent
        self.parent.geometry('620x600')
        self.parent.configure(bg=COLORS['bg'])
        self.countdown_running = False
        self.countdown_thread = None
        self.remaining = 0
        self.setup_ui()

    def setup_ui(self):
        # Title bar
        title = tk.Label(self.parent, text='⏻  Auto Shutdown / Scheduler',
                          font=FONTS['large'], bg=COLORS['bg_header'],
                          fg=COLORS['white'], pady=10)
        title.pack(fill='x')

        main = tk.Frame(self.parent, bg=COLORS['bg'])
        main.pack(fill='both', expand=True, padx=15, pady=10)

        # ── Action Section ──────────────────────────────────────────────
        action_frame = tk.LabelFrame(main, text='  Action  ',
                                      font=FONTS['subtitle'],
                                      bg=COLORS['bg'], fg=COLORS['text'],
                                      relief='groove', bd=2)
        action_frame.pack(fill='x', pady=(0, 10))

        self.action_var = tk.StringVar(value='shutdown')
        actions = [
            ('⏻ Shutdown',   'shutdown'),
            ('🔄 Restart',    'restart'),
            ('💤 Hibernate',  'hibernate'),
            ('🚪 Logoff',     'logoff'),
            ('😴 Standby',    'standby'),
            ('🔒 Lock',       'lock'),
            ('❌ Close App',  'close_app'),
            ('▶️ Run App',   'run_app'),
            ('💬 Message',   'message'),
        ]
        row, col = 0, 0
        for text, val in actions:
            rb = tk.Radiobutton(action_frame, text=text, value=val,
                                 variable=self.action_var,
                                 font=FONTS['normal'],
                                 bg=COLORS['bg'], fg=COLORS['text'],
                                 selectcolor=COLORS['selected'],
                                 command=self.on_action_change,
                                 cursor='hand2')
            rb.grid(row=row, column=col, sticky='w', padx=10, pady=3)
            col += 1
            if col >= 3:
                col = 0
                row += 1

        # ── App / Message Config ────────────────────────────────────────
        self.extra_frame = tk.LabelFrame(main, text='  Configuration  ',
                                          font=FONTS['subtitle'],
                                          bg=COLORS['bg'], fg=COLORS['text'],
                                          relief='groove', bd=2)
        self.extra_frame.pack(fill='x', pady=(0, 10))

        tk.Label(self.extra_frame, text='App/Message:',
                  font=FONTS['normal'], bg=COLORS['bg']).grid(
            row=0, column=0, padx=8, pady=5, sticky='w')
        self.app_entry = tk.Entry(self.extra_frame, font=FONTS['normal'], width=40)
        self.app_entry.grid(row=0, column=1, padx=5, pady=5, sticky='ew')
        tk.Button(self.extra_frame, text='Browse...', font=FONTS['small'],
                   bg=COLORS['info'], fg='white', relief='flat',
                   cursor='hand2',
                   command=self.browse_app).grid(row=0, column=2, padx=5, pady=5)
        self.extra_frame.columnconfigure(1, weight=1)

        # ── Timer Section ───────────────────────────────────────────────
        timer_frame = tk.LabelFrame(main, text='  Timer  ',
                                     font=FONTS['subtitle'],
                                     bg=COLORS['bg'], fg=COLORS['text'],
                                     relief='groove', bd=2)
        timer_frame.pack(fill='x', pady=(0, 10))

        inner = tk.Frame(timer_frame, bg=COLORS['bg'])
        inner.pack(padx=10, pady=8)

        # Mode: immediate vs scheduled
        self.timer_mode = tk.StringVar(value='countdown')
        tk.Radiobutton(inner, text='Ngay lập tức', value='immediate',
                        variable=self.timer_mode, font=FONTS['normal'],
                        bg=COLORS['bg'], fg=COLORS['text'],
                        selectcolor=COLORS['selected'],
                        command=self.toggle_timer).pack(side='left', padx=5)
        tk.Radiobutton(inner, text='Đếm ngược (phút):', value='countdown',
                        variable=self.timer_mode, font=FONTS['normal'],
                        bg=COLORS['bg'], fg=COLORS['text'],
                        selectcolor=COLORS['selected'],
                        command=self.toggle_timer).pack(side='left', padx=5)

        self.minutes_var = tk.StringVar(value='30')
        self.min_entry = tk.Spinbox(inner, from_=1, to=9999,
                                     textvariable=self.minutes_var,
                                     font=FONTS['normal'], width=6)
        self.min_entry.pack(side='left', padx=5)

        tk.Label(inner, text='phút', font=FONTS['normal'],
                  bg=COLORS['bg']).pack(side='left')

        # Specific time
        time_row = tk.Frame(timer_frame, bg=COLORS['bg'])
        time_row.pack(padx=10, pady=(0, 8))
        tk.Radiobutton(time_row, text='Lúc giờ cụ thể:', value='specific',
                        variable=self.timer_mode, font=FONTS['normal'],
                        bg=COLORS['bg'], fg=COLORS['text'],
                        selectcolor=COLORS['selected'],
                        command=self.toggle_timer).pack(side='left', padx=5)
        self.time_entry = tk.Entry(time_row, font=FONTS['normal'], width=8)
        self.time_entry.insert(0, '23:00')
        self.time_entry.pack(side='left', padx=5)
        tk.Label(time_row, text='(HH:MM)', font=FONTS['small'],
                  bg=COLORS['bg'], fg=COLORS['text_light']).pack(side='left')

        # ── Countdown Display ───────────────────────────────────────────
        self.countdown_label = tk.Label(main,
                                         text='⏱ 00:00:00',
                                         font=('Segoe UI', 22, 'bold'),
                                         bg=COLORS['bg'], fg=COLORS['accent'])
        self.countdown_label.pack(pady=5)

        # ── Control Buttons ─────────────────────────────────────────────
        btn_frame = tk.Frame(main, bg=COLORS['bg'])
        btn_frame.pack(pady=5)

        self.btn_start = tk.Button(btn_frame, text='▶ Start',
                                    font=FONTS['subtitle'],
                                    bg=COLORS['accent'], fg='white',
                                    relief='flat', width=12, height=2,
                                    cursor='hand2',
                                    command=self.start_action)
        self.btn_start.pack(side='left', padx=8)

        self.btn_cancel = tk.Button(btn_frame, text='⏹ Cancel',
                                     font=FONTS['subtitle'],
                                     bg=COLORS['danger'], fg='white',
                                     relief='flat', width=12, height=2,
                                     cursor='hand2',
                                     state='disabled',
                                     command=self.cancel_action)
        self.btn_cancel.pack(side='left', padx=8)

        btn_now = tk.Button(btn_frame, text='⚡ Ngay Bây Giờ',
                             font=FONTS['subtitle'],
                             bg=COLORS['warning'], fg='white',
                             relief='flat', width=14, height=2,
                             cursor='hand2',
                             command=self.do_action_now)
        btn_now.pack(side='left', padx=8)

        btn_logs = tk.Button(btn_frame, text='📜 Xem Lịch Sử Log',
                             font=FONTS['subtitle'],
                             bg=COLORS['info'], fg='white',
                             relief='flat', width=16, height=2,
                             cursor='hand2',
                             command=self.open_log_history_window)
        btn_logs.pack(side='left', padx=8)

    def on_action_change(self):
        pass

    def browse_app(self):
        path = filedialog.askopenfilename(
            title='Select Application',
            filetypes=[('Executables', '*.exe'), ('All files', '*.*')]
        )
        if path:
            self.app_entry.delete(0, 'end')
            self.app_entry.insert(0, path)

    def toggle_timer(self):
        mode = self.timer_mode.get()
        self.min_entry.configure(state='normal' if mode == 'countdown' else 'disabled')
        self.time_entry.configure(state='normal' if mode == 'specific' else 'disabled')

    def start_action(self):
        if self.countdown_running:
            return
        mode = self.timer_mode.get()
        if mode == 'immediate':
            self.do_action_now()
            return
        elif mode == 'countdown':
            try:
                mins = int(self.minutes_var.get())
                self.remaining = mins * 60
            except:
                messagebox.showerror('Error', 'Số phút không hợp lệ!')
                return
        elif mode == 'specific':
            import datetime
            try:
                t = self.time_entry.get().strip()
                h, m = map(int, t.split(':'))
                now = datetime.datetime.now()
                target = now.replace(hour=h, minute=m, second=0, microsecond=0)
                if target <= now:
                    target += datetime.timedelta(days=1)
                self.remaining = int((target - now).total_seconds())
            except:
                messagebox.showerror('Error', 'Định dạng giờ không hợp lệ! Dùng HH:MM')
                return

        self.countdown_running = True
        self.btn_start.configure(state='disabled')
        self.btn_cancel.configure(state='normal')
        self.countdown_thread = threading.Thread(target=self.run_countdown, daemon=True)
        self.countdown_thread.start()

    def run_countdown(self):
        while self.remaining > 0 and self.countdown_running:
            h = self.remaining // 3600
            m = (self.remaining % 3600) // 60
            s = self.remaining % 60
            self.countdown_label.configure(
                text=f'⏱ {h:02d}:{m:02d}:{s:02d}',
                fg=COLORS['danger'] if self.remaining <= 60 else COLORS['accent']
            )
            time.sleep(1)
            self.remaining -= 1

        if self.countdown_running:
            self.do_action_now()
        self.countdown_running = False
        self.btn_start.configure(state='normal')
        self.btn_cancel.configure(state='disabled')
        self.countdown_label.configure(text='⏱ 00:00:00', fg=COLORS['accent'])

    def cancel_action(self):
        self.countdown_running = False
        self.btn_start.configure(state='normal')
        self.btn_cancel.configure(state='disabled')
        self.countdown_label.configure(text='⏱ 00:00:00', fg=COLORS['accent'])
        # Cancel Windows shutdown if scheduled
        subprocess.run(['shutdown', '/a'], capture_output=True)

    def do_action_now(self):
        action = self.action_var.get()
        try:
            if action == 'shutdown':
                subprocess.run(['shutdown', '/s', '/t', '0'])
            elif action == 'restart':
                subprocess.run(['shutdown', '/r', '/t', '0'])
            elif action == 'hibernate':
                subprocess.run(['shutdown', '/h'])
            elif action == 'logoff':
                subprocess.run(['shutdown', '/l'])
            elif action == 'standby':
                # Use SetSuspendState via powershell
                subprocess.run(['powershell', '-Command',
                                 'Add-Type -Assembly System.Windows.Forms; '
                                 '[System.Windows.Forms.Application]::SetSuspendState("Suspend",$false,$false)'])
            elif action == 'lock':
                import ctypes
                ctypes.windll.user32.LockWorkStation()
            elif action == 'close_app':
                app = self.app_entry.get().strip()
                if app:
                    subprocess.run(['taskkill', '/f', '/im', os.path.basename(app)])
                else:
                    messagebox.showwarning('Warning', 'Nhập tên ứng dụng cần đóng!')
            elif action == 'run_app':
                app = self.app_entry.get().strip()
                if app:
                    subprocess.Popen(app)
                else:
                    messagebox.showwarning('Warning', 'Chọn ứng dụng cần chạy!')
            elif action == 'message':
                msg = self.app_entry.get().strip() or 'IT Tool LTT Notification'
                messagebox.showinfo('Thông báo', msg)
        except Exception as e:
            messagebox.showerror('Error', f'Lỗi thực thi: {e}')

    def open_log_history_window(self):
        """Opens Tkinter window showing shutdown and reboot event logs."""
        win = tk.Toplevel(self.parent)
        win.title('📜 Lịch Sử Tắt Máy & Sập Nguồn (Windows System Event Logs)')
        win.geometry('900x560')
        win.configure(bg=COLORS['bg'])

        header = tk.Frame(win, bg=COLORS['bg_header'], pady=10, padx=15)
        header.pack(fill='x')
        tk.Label(header, text='📜 Lịch Sử Tắt Máy, Khởi Động & Sập Nguồn (Event ID 1074, 6008, 41)',
                 font=FONTS['subtitle'], bg=COLORS['bg_header'], fg=COLORS['white']).pack(side='left')

        btn_bar = tk.Frame(win, bg=COLORS['bg'], padx=10, pady=8)
        btn_bar.pack(fill='x')

        status_lbl = tk.Label(btn_bar, text='Đang tải dữ liệu thời gian thực...', font=FONTS['small'], bg=COLORS['bg'], fg=COLORS['text_light'])
        status_lbl.pack(side='left')

        # Treeview
        tree_frame = tk.Frame(win, bg=COLORS['bg'], padx=10)
        tree_frame.pack(fill='both', expand=True)

        cols = ('time', 'id', 'action', 'user', 'process', 'reason')
        tree = ttk.Treeview(tree_frame, columns=cols, show='headings', selectmode='browse', height=12)
        tree.heading('time', text='Thời Gian (Real-Time)')
        tree.heading('id', text='Event ID')
        tree.heading('action', text='Hành Động')
        tree.heading('user', text='Người Dùng')
        tree.heading('process', text='Tiến Trình')
        tree.heading('reason', text='Lý Do / Diễn Giải')

        tree.column('time', width=160, anchor='center')
        tree.column('id', width=75, anchor='center')
        tree.column('action', width=140, anchor='w')
        tree.column('user', width=130, anchor='w')
        tree.column('process', width=120, anchor='w')
        tree.column('reason', width=240, anchor='w')

        sb = ttk.Scrollbar(tree_frame, orient='vertical', command=tree.yview)
        tree.configure(yscrollcommand=sb.set)
        tree.pack(side='left', fill='both', expand=True)
        sb.pack(side='right', fill='y')

        # Details box
        det_frame = tk.LabelFrame(win, text='  Chi Tiết Bản Ghi  ', font=FONTS['small'], bg=COLORS['bg'], fg=COLORS['text'])
        det_frame.pack(fill='x', padx=10, pady=8)
        det_txt = tk.Text(det_frame, height=5, font=('Consolas', 9), wrap='word', bg='#f8fafc')
        det_txt.pack(fill='both', expand=True, padx=5, pady=5)

        current_logs = []

        def load_logs():
            nonlocal current_logs
            status_lbl.configure(text='⏳ Đang đọc nhật ký Windows System Event Log...')
            tree.delete(*tree.get_children())
            det_txt.delete('1.0', 'end')
            res = get_shutdown_event_logs(max_events=25)
            if res.get('success'):
                current_logs = res.get('logs', [])
                for idx, item in enumerate(current_logs):
                    tree.insert('', 'end', iid=str(idx), values=(
                        item.get('friendly_time', item.get('time')),
                        item.get('id'),
                        item.get('action_label'),
                        item.get('user') or 'System',
                        item.get('process_name') or 'N/A',
                        item.get('summary')
                    ))
                stats = res.get('stats', {})
                status_lbl.configure(text=f"🟢 Đã nạp {len(current_logs)} bản ghi | Lần gần nhất: {stats.get('last_event_time', 'N/A')} | Cập nhật lúc {stats.get('query_time', '')}")
            else:
                status_lbl.configure(text=f"❌ Lỗi: {res.get('message')}")

        def on_select(event):
            sel = tree.selection()
            if sel:
                idx = int(sel[0])
                if idx < len(current_logs):
                    item = current_logs[idx]
                    det_txt.delete('1.0', 'end')
                    det = f"ID: {item.get('id')} | Mức độ: {item.get('level')} | Thời gian: {item.get('friendly_time')} ({item.get('relative_time')})\n"
                    det += f"Người dùng: {item.get('user')} | Máy tính: {item.get('machine')}\n"
                    det += f"Tiến trình: {item.get('process')}\n"
                    det += f"Lý do: {item.get('reason')} (Reason Code: {item.get('reason_code')})\n"
                    det += f"Ghi chú IT: {item.get('expert_note')}\n"
                    det += f"Raw Message: {item.get('raw_message')}"
                    det_txt.insert('1.0', det)

        tree.bind('<<TreeviewSelect>>', on_select)

        tk.Button(btn_bar, text='🔄 Làm Mới (Real-time)', font=FONTS['small'], bg=COLORS['accent'], fg='white', relief='flat', command=lambda: threading.Thread(target=load_logs, daemon=True).start()).pack(side='right', padx=4)
        tk.Button(btn_bar, text='📊 Xuất Báo Cáo Excel', font=FONTS['small'], bg=COLORS['success'], fg='white', relief='flat', command=lambda: export_shutdown_event_logs(100, format_type='excel')).pack(side='right', padx=4)

        threading.Thread(target=load_logs, daemon=True).start()


# ── STANDALONE TASK SCHEDULER BACKEND API FUNCTIONS ───────────────────────

import base64

def _run_ps(ps_script, timeout=10):
    """Executes PowerShell script via Base64 EncodedCommand to avoid escaping issues."""
    encoded = base64.b64encode(ps_script.encode('utf-16le')).decode('utf-8')
    return subprocess.run(["powershell", "-NoProfile", "-EncodedCommand", encoded], capture_output=True, text=True, timeout=timeout)


def execute_immediate_power(action):
    """Executes immediate system power operation (Shutdown, Restart, Sleep, Logoff, Lock)."""
    try:
        if action in ("shutdown", "shutdown_now"):
            cmd = "shutdown /s /f /t 0"
            msg = "Đã phát lệnh Tắt máy ngay lập tức."
        elif action in ("restart", "restart_now"):
            cmd = "shutdown /r /f /t 0"
            msg = "Đã phát lệnh Khởi động lại ngay lập tức."
        elif action in ("hibernate", "sleep"):
            cmd = 'powershell -Command "Add-Type -Assembly System.Windows.Forms; [System.Windows.Forms.Application]::SetSuspendState(\'Suspend\', $false, $false)"'
            msg = "Đã chuyển hệ thống sang chế độ Sleep/Hibernate."
        elif action == "logoff":
            cmd = "shutdown /l"
            msg = "Đã đăng xuất tài khoản người dùng."
        elif action == "lock":
            cmd = "rundll32.exe user32.dll,LockWorkStation"
            msg = "Đã khóa màn hình máy tính."
        elif action == "cancel":
            cmd = "shutdown /a"
            msg = "Đã hủy lệnh hẹn giờ tắt máy hệ thống."
        else:
            return {"success": False, "message": f"Hành động không hợp lệ: {action}"}

        subprocess.run(cmd, shell=True)
        return {"success": True, "message": msg}
    except Exception as e:
        return {"success": False, "message": f"Lỗi thực thi: {str(e)}"}


def create_quick_timer_task(action, minutes, message=""):
    """Creates a one-time Windows Task Scheduler task for quick timer, ensuring execution even if app is closed."""
    import datetime
    try:
        sec = float(minutes) * 60
        if sec <= 0:
            return {"success": False, "message": "Số phút phải lớn hơn 0!"}
    except Exception as e:
        return {"success": False, "message": f"Số phút không hợp lệ: {e}"}

    target_dt = datetime.datetime.now() + datetime.timedelta(seconds=sec)
    dt_str = target_dt.strftime("%Y-%m-%d %H:%M:%S")

    # Sanitize message to prevent PowerShell injection
    safe_msg = (message.strip() or f"Tu dong {action} boi IT Tool LTT Quick Timer").replace('"', '').replace("'", "").replace("`", "")

    if action == "shutdown":
        exe = "shutdown.exe"
        args = f'/s /f /t 0 /c "{safe_msg}"'
    elif action == "restart":
        exe = "shutdown.exe"
        args = f'/r /f /t 0 /c "{safe_msg}"'
    elif action == "hibernate":
        exe = "powershell.exe"
        args = '-Command "Add-Type -Assembly System.Windows.Forms; [System.Windows.Forms.Application]::SetSuspendState(\'Suspend\', $false, $false)"'
    elif action == "logoff":
        exe = "shutdown.exe"
        args = "/l"
    elif action == "lock":
        exe = "%windir%\\System32\\rundll32.exe"
        args = "user32.dll,LockWorkStation"
    else:
        exe = "shutdown.exe"
        args = f'/s /f /t 0 /c "{safe_msg}"'

    full_name = "ITTools_Shutdown_QuickTimer"

    ps_code = f"""
    $dt = [datetime]::ParseExact('{dt_str}', 'yyyy-MM-dd HH:mm:ss', $null)
    $trig = New-ScheduledTaskTrigger -Once -At $dt
    $act = New-ScheduledTaskAction -Execute "{exe}" -Argument '{args}'
    $princ = New-ScheduledTaskPrincipal -UserId "$env:USERNAME" -LogonType Interactive -RunLevel Highest
    $settings = New-ScheduledTaskSettingsSet -ExecutionTimeLimit (New-TimeSpan -Minutes 5) -AllowStartIfOnBatteries -DontStopIfGoingOnBatteries
    Register-ScheduledTask -TaskName "{full_name}" -Trigger $trig -Action $act -Principal $princ -Settings $settings -Force
    """

    try:
        r = _run_ps(ps_code, timeout=10)
        if r.returncode == 0:
            action_names = {
                "shutdown": "Tắt máy",
                "restart": "Khởi động lại",
                "hibernate": "Sleep / Hibernate",
                "logoff": "Đăng xuất",
                "lock": "Khóa màn hình"
            }
            act_label = action_names.get(action, action)
            return {
                "success": True,
                "message": f"Đã lưu lịch {act_label} lúc {target_dt.strftime('%H:%M:%S')} vào Windows Task Scheduler. Dù bạn TẮT PHẦN MỀM thì máy vẫn sẽ tự động thực thi!",
                "target_time": dt_str
            }
        else:
            return {"success": False, "message": f"Lỗi tạo lịch Windows Task Scheduler: {r.stderr or r.stdout}"}
    except Exception as e:
        return {"success": False, "message": f"Lỗi ngoại lệ: {str(e)}"}


def cancel_quick_timer_task():
    """Cancels the quick timer Windows Task Scheduler task."""
    full_name = "ITTools_Shutdown_QuickTimer"
    ps_code = f'Unregister-ScheduledTask -TaskName "{full_name}" -Confirm:$false -ErrorAction SilentlyContinue'
    try:
        _run_ps(ps_code, timeout=8)
        subprocess.run("shutdown /a", shell=True, capture_output=True)
        return {"success": True, "message": "Đã hủy lịch đếm ngược Quick Timer trong Windows Task Scheduler!"}
    except Exception as e:
        return {"success": False, "message": str(e)}


def get_scheduled_shutdown_tasks():
    """Returns all scheduled shutdown tasks created by IT Tool LTT."""
    ps_code = """
    $tasks = Get-ScheduledTask -TaskPath '\\' -ErrorAction SilentlyContinue | Where-Object { $_.TaskName -like 'ITTools_Shutdown_*' }
    if (-not $tasks) { Write-Output "[]"; exit }
    $res = @()
    foreach ($t in $tasks) {
        $info = Get-ScheduledTaskInfo -TaskName $t.TaskName -ErrorAction SilentlyContinue
        $res += [PSCustomObject]@{
            TaskName = $t.TaskName
            State = $t.State.ToString()
            NextRun = if ($info.NextRunTime) { $info.NextRunTime.ToString('yyyy-MM-dd HH:mm:ss') } else { '' }
            LastRun = if ($info.LastRunTime) { $info.LastRunTime.ToString('yyyy-MM-dd HH:mm:ss') } else { '' }
            Execute = $t.Actions.Execute
            Arguments = $t.Actions.Arguments
        }
    }
    @($res) | ConvertTo-Json -Depth 3
    """
    try:
        r = _run_ps(ps_code, timeout=10)
        if r.returncode == 0 and r.stdout.strip():
            data = json.loads(r.stdout)
            if not data:
                return {"success": True, "tasks": []}
            if isinstance(data, dict):
                data = [data]
            if not isinstance(data, list):
                return {"success": True, "tasks": []}
            parsed = []
            for item in data:
                if not isinstance(item, dict):
                    continue
                tn = item.get("TaskName", "") or ""
                disp_name = tn.replace("ITTools_Shutdown_", "").replace("_", " ")
                args = item.get("Arguments", "") or ""
                
                action_type = "shutdown"
                if "/r" in args:
                    action_type = "restart"
                elif "/h" in args:
                    action_type = "hibernate"
                elif "/l" in args:
                    action_type = "logoff"
                elif "LockWorkStation" in args or "user32.dll" in args:
                    action_type = "lock"

                parsed.append({
                    "task_name": tn,
                    "display_name": disp_name,
                    "action": action_type,
                    "state": item.get("State", "Unknown"),
                    "enabled": (item.get("State", "") or "").lower() != "disabled",
                    "next_run": item.get("NextRun", "") or "N/A",
                    "last_run": item.get("LastRun", "") or "N/A",
                    "arguments": args
                })
            return {"success": True, "tasks": parsed}
    except Exception as e:
        print("Lỗi get_scheduled_shutdown_tasks:", e)
    return {"success": True, "tasks": []}


def create_scheduled_shutdown_task(name, action, freq, time_str, days=None, date_str="", message=""):
    """Creates a Windows Task Scheduler auto shutdown task."""
    import datetime as dt_module
    if not name or not name.strip():
        return {"success": False, "message": "Vui lòng nhập tên lịch hẹn giờ!"}
    
    clean_name = "".join(c for c in name.strip() if c.isalnum() or c in ("_", "-")).strip()
    if not clean_name:
        clean_name = "Task_" + dt_module.datetime.now().strftime("%Y%m%d%H%M%S")
    full_name = f"ITTools_Shutdown_{clean_name}"
    
    # Sanitize message to prevent PowerShell injection
    safe_msg = (message.strip() or f"Tu dong {action} boi IT Tool LTT").replace('"', '').replace("'", "").replace("`", "")

    if action == "shutdown":
        exe = "shutdown.exe"
        args = f'/s /f /t 60 /c "{safe_msg}"'
    elif action == "restart":
        exe = "shutdown.exe"
        args = f'/r /f /t 60 /c "{safe_msg}"'
    elif action == "hibernate":
        exe = "shutdown.exe"
        args = "/h"
    elif action == "logoff":
        exe = "shutdown.exe"
        args = "/l"
    elif action == "lock":
        exe = "%windir%\\System32\\rundll32.exe"
        args = "user32.dll,LockWorkStation"
    else:
        exe = "shutdown.exe"
        args = f'/s /f /t 60 /c "{safe_msg}"'

    time_val = time_str.strip() or "23:00"

    if freq == "daily":
        trig_cmd = f'New-ScheduledTaskTrigger -Daily -At "{time_val}"'
    elif freq == "weekly":
        if not days or len(days) == 0:
            days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
        days_str = ",".join(days)
        trig_cmd = f'New-ScheduledTaskTrigger -Weekly -DaysOfWeek {days_str} -At "{time_val}"'
    elif freq == "once":
        today_str = dt_module.datetime.now().strftime('%Y-%m-%d')
        dt_val = f"{date_str.strip()} {time_val}" if date_str.strip() else f"{today_str} {time_val}"
        trig_cmd = f'New-ScheduledTaskTrigger -Once -At "{dt_val}"'
    else:
        trig_cmd = f'New-ScheduledTaskTrigger -Daily -At "{time_val}"'

    ps_code = f"""
    $trig = {trig_cmd}
    $act = New-ScheduledTaskAction -Execute "{exe}" -Argument '{args}'
    $princ = New-ScheduledTaskPrincipal -UserId "$env:USERNAME" -LogonType Interactive -RunLevel Highest
    $settings = New-ScheduledTaskSettingsSet -ExecutionTimeLimit (New-TimeSpan -Minutes 10) -AllowStartIfOnBatteries -DontStopIfGoingOnBatteries -WakeToRun
    Register-ScheduledTask -TaskName "{full_name}" -Trigger $trig -Action $act -Principal $princ -Settings $settings -Force
    """
    try:
        r = _run_ps(ps_code, timeout=10)
        if r.returncode == 0:
            return {"success": True, "message": f"Đã tạo thành công lịch '{clean_name}' ({freq.upper()} lúc {time_val})!"}
        else:
            return {"success": False, "message": f"Lỗi PowerShell: {r.stderr or r.stdout}"}
    except Exception as e:
        return {"success": False, "message": f"Lỗi ngoại lệ: {str(e)}"}


def toggle_scheduled_shutdown_task(task_name, enable):
    """Enables or disables a scheduled shutdown task."""
    cmd_verb = "Enable-ScheduledTask" if enable else "Disable-ScheduledTask"
    ps_code = f'{cmd_verb} -TaskName "{task_name}"'
    try:
        r = _run_ps(ps_code, timeout=8)
        if r.returncode == 0:
            st = "bật" if enable else "tắt"
            return {"success": True, "message": f"Đã {st} lịch hẹn giờ '{task_name.replace('ITTools_Shutdown_', '')}' thành công!"}
        else:
            return {"success": False, "message": r.stderr or r.stdout}
    except Exception as e:
        return {"success": False, "message": str(e)}


def delete_scheduled_shutdown_task(task_name):
    """Deletes a scheduled shutdown task."""
    ps_code = f'Unregister-ScheduledTask -TaskName "{task_name}" -Confirm:$false'
    try:
        r = _run_ps(ps_code, timeout=8)
        if r.returncode == 0:
            return {"success": True, "message": f"Đã xóa lịch hẹn giờ '{task_name.replace('ITTools_Shutdown_', '')}'!"}
        else:
            return {"success": False, "message": r.stderr or r.stdout}
    except Exception as e:
        return {"success": False, "message": str(e)}


def run_scheduled_shutdown_task_now(task_name):
    """Executes a scheduled shutdown task immediately."""
    ps_code = f'Start-ScheduledTask -TaskName "{task_name}"'
    try:
        r = _run_ps(ps_code, timeout=8)
        if r.returncode == 0:
            return {"success": True, "message": f"Đã kích hoạt chạy ngay lịch '{task_name.replace('ITTools_Shutdown_', '')}'!"}
        else:
            return {"success": False, "message": r.stderr or r.stdout}
    except Exception as e:
        return {"success": False, "message": str(e)}


def get_shutdown_event_logs(max_events=10, event_ids="1074,6008,41"):
    """
    Fetches Windows System Event Logs for Shutdown, Reboot, Unexpected Shutdown, and Kernel-Power events.
    Ids:
      - 1074: Clean Shutdown/Restart initiated by User or Process
      - 6008: Unexpected Shutdown (Mất nguồn điện, tắt nóng)
      - 41: Kernel-Power (System rebooted without cleanly shutting down first, BSOD, crash)
    """
    try:
        max_events = int(max_events)
        if max_events <= 0:
            max_events = 10
        elif max_events > 500:
            max_events = 500
    except Exception:
        max_events = 10

    clean_ids = ",".join([i.strip() for i in str(event_ids).split(",") if i.strip().isdigit()])
    if not clean_ids:
        clean_ids = "1074,6008,41"

    ps_code = f"""
    $events = Get-WinEvent -FilterHashtable @{{LogName='System'; Id={clean_ids}}} -MaxEvents {max_events} -ErrorAction SilentlyContinue
    if (-not $events) {{ Write-Output '[]'; exit }}
    $arr = @()
    foreach ($e in $events) {{
        $arr += [PSCustomObject]@{{
            Id = $e.Id
            TimeCreated = if ($e.TimeCreated) {{ $e.TimeCreated.ToString('yyyy-MM-dd HH:mm:ss') }} else {{ '' }}
            ProviderName = $e.ProviderName
            Level = $e.LevelDisplayName
            MachineName = $e.MachineName
            Message = $e.Message
        }}
    }}
    @($arr) | ConvertTo-Json -Depth 2
    """
    try:
        r = _run_ps(ps_code, timeout=15)
        raw = json.loads(r.stdout) if (r.returncode == 0 and r.stdout.strip()) else []
        if isinstance(raw, dict):
            raw = [raw]
    except Exception as e:
        return {"success": False, "message": str(e), "logs": [], "stats": {}}

    now = datetime.datetime.now()
    parsed = []
    for e in raw:
        if not isinstance(e, dict):
            continue
        msg = e.get('Message', '') or ''
        eid = e.get('Id')
        time_str = e.get('TimeCreated', '') or ''
        machine = e.get('MachineName', '') or ''

        # Calculate relative time and friendly date
        rel_time = time_str
        friendly_date = time_str
        try:
            dt = datetime.datetime.strptime(time_str, '%Y-%m-%d %H:%M:%S')
            friendly_date = dt.strftime('%d/%m/%Y %H:%M:%S')
            diff = now - dt
            seconds = int(diff.total_seconds())
            if seconds < 60:
                rel_time = f'{max(0, seconds)} giây trước'
            elif seconds < 3600:
                rel_time = f'{seconds // 60} phút trước'
            elif seconds < 86400:
                rel_time = f'{seconds // 3600} giờ trước'
            elif seconds < 86400 * 7:
                rel_time = f'{seconds // 86400} ngày trước'
            else:
                rel_time = dt.strftime('%d/%m/%Y %H:%M')
        except Exception:
            pass

        item = {
            'id': eid,
            'time': time_str,
            'friendly_time': friendly_date,
            'relative_time': rel_time,
            'provider': e.get('ProviderName') or '',
            'machine': machine,
            'level': e.get('Level') or ('Critical' if eid == 41 else ('Warning' if eid == 6008 else 'Information')),
            'raw_message': msg,
            'action_type': 'unknown',
            'action_label': 'Không rõ',
            'process': '',
            'process_name': '',
            'user': '',
            'reason': '',
            'reason_code': '',
            'shutdown_type': '',
            'comment': '',
            'expert_note': '',
            'summary': ''
        }

        if eid == 1074:
            m_proc = re.search(r'The process ([^\(]+?)(?:\s+\([^\)]+\))?\s+has initiated the (\w[\w\s]*?) of computer', msg, re.IGNORECASE)
            if m_proc:
                item['process'] = m_proc.group(1).strip()
                action_word = m_proc.group(2).strip().lower()
                if 'restart' in action_word:
                    item['action_type'] = 'restart'
                    item['action_label'] = 'Khởi động lại (Restart)'
                    item['shutdown_type'] = 'restart'
                elif 'power off' in action_word:
                    item['action_type'] = 'poweroff'
                    item['action_label'] = 'Tắt nguồn (Power Off)'
                    item['shutdown_type'] = 'power off'
                else:
                    item['action_type'] = 'shutdown'
                    item['action_label'] = 'Tắt máy (Shutdown)'
                    item['shutdown_type'] = 'shutdown'
            else:
                if 'restart' in msg.lower():
                    item['action_type'] = 'restart'
                    item['action_label'] = 'Khởi động lại (Restart)'
                    item['shutdown_type'] = 'restart'
                elif 'power off' in msg.lower():
                    item['action_type'] = 'poweroff'
                    item['action_label'] = 'Tắt nguồn (Power Off)'
                    item['shutdown_type'] = 'power off'
                else:
                    item['action_type'] = 'shutdown'
                    item['action_label'] = 'Tắt máy (Shutdown)'
                    item['shutdown_type'] = 'shutdown'

            m_user = re.search(r'on behalf of user ([^\r\n]+?) for the following reason', msg, re.IGNORECASE)
            if m_user:
                item['user'] = m_user.group(1).strip()

            m_reason = re.search(r'for the following reason:\s*([^\r\n]+)', msg, re.IGNORECASE)
            if m_reason:
                item['reason'] = m_reason.group(1).strip()

            m_code = re.search(r'Reason Code:\s*(0x[0-9a-fA-F]+)', msg, re.IGNORECASE)
            if m_code:
                item['reason_code'] = m_code.group(1).strip()

            m_comment = re.search(r'Comment:\s*([^\r\n]*)', msg, re.IGNORECASE)
            if m_comment:
                item['comment'] = m_comment.group(1).strip()

            proc_name = item['process'].split('\\')[-1] if item['process'] else ''
            item['process_name'] = proc_name
            user_name = item['user'] or 'User'
            reason_desc = item['reason'] or 'Bình thường'
            item['summary'] = f"{user_name} ({proc_name or 'Hệ thống'}) - {reason_desc}"
            item['expert_note'] = f"Lệnh {item['action_label']} chủ động, phát lệnh bởi tiến trình {proc_name or 'System'} dưới quyền người dùng {user_name}."

        elif eid == 6008:
            item['action_type'] = 'unexpected_shutdown'
            item['action_label'] = 'Sập nguồn đột ngột'
            item['reason'] = 'Tắt máy không mong muốn (Dirty Shutdown)'
            m_prev = re.search(r'previous system shutdown at ([^\r\n]+) was unexpected', msg, re.IGNORECASE)
            if m_prev:
                item['reason'] = f"Thời điểm sập nguồn: {m_prev.group(1).strip()}"
            item['summary'] = 'Hệ thống bị ngắt nguồn đột ngột (mất điện, rút phích cắm hoặc bấm giữ nút nguồn)'
            item['expert_note'] = 'Hệ điều hành không kịp ghi nhật ký tắt bình thường. Nguyên nhân thường do cúp điện lưới, hỏng bộ nguồn PSU/Pin hoặc người dùng ép tắt cưỡng bức.'

        elif eid == 41:
            item['action_type'] = 'kernel_power'
            item['action_label'] = 'Kernel-Power (Reboot không an toàn)'
            item['reason'] = 'Reboot không qua quy trình tắt sạch'
            item['summary'] = 'Hệ thống khởi động lại sau khi bị mất nguồn hoặc lỗi phần cứng / BSOD'
            item['expert_note'] = 'Windows khởi động lại sau khi bị treo đơ, sập nguồn hoặc màn hình xanh (BSOD). Cần kiểm tra nhiệt độ CPU, RAM, Driver hoặc nguồn điện.'

        parsed.append(item)

    stats = {
        'total': len(parsed),
        'shutdown_count': sum(1 for x in parsed if x['action_type'] in ('shutdown', 'poweroff')),
        'restart_count': sum(1 for x in parsed if x['action_type'] == 'restart'),
        'unexpected_count': sum(1 for x in parsed if x['id'] == 6008),
        'kernel_power_count': sum(1 for x in parsed if x['id'] == 41),
        'last_event_time': parsed[0]['friendly_time'] if parsed else 'Chưa có',
        'last_event_label': parsed[0]['action_label'] if parsed else 'Chưa có',
        'query_time': now.strftime('%H:%M:%S - %d/%m/%Y')
    }

    return {'success': True, 'logs': parsed, 'stats': stats}


def export_shutdown_event_logs(max_events=100, event_ids="1074,6008,41", save_path="", format_type="excel"):
    """Exports shutdown event logs to a beautifully formatted Excel (.xlsx) or CSV file."""
    res = get_shutdown_event_logs(max_events, event_ids)
    if not res.get("success"):
        return res

    logs = res.get("logs", [])
    stats = res.get("stats", {})

    user_profile = os.environ.get('USERPROFILE', os.path.expanduser('~'))
    desktop = os.path.join(user_profile, 'Desktop')
    if not os.path.exists(desktop):
        desktop = user_profile

    ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    comp_name = logs[0].get('machine') if logs else os.environ.get('COMPUTERNAME', 'LOCAL_PC')

    # Excel export via openpyxl
    if format_type.lower() == 'excel' or (save_path and save_path.lower().endswith('.xlsx')):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            if not save_path:
                save_path = os.path.join(desktop, f"Lich_Su_Tat_May_{ts}.xlsx")

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Lịch Sử Tắt Máy"
            ws.views.sheetView[0].showGridLines = True

            # 1. Main Title Banner (Row 1)
            ws.merge_cells("A1:L1")
            title_cell = ws["A1"]
            title_cell.value = "BÁO CÁO CHI TIẾT LỊCH SỬ TẮT MÁY & KHỞI ĐỘNG HỆ THỐNG"
            title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
            title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 36

            # 2. Subtitle / Metadata Banner (Row 2)
            ws.merge_cells("A2:L2")
            sub_cell = ws["A2"]
            now_str = datetime.datetime.now().strftime('%d/%m/%Y %H:%M:%S')
            sub_cell.value = f"🖥️ Máy tính: {comp_name}  |  ⏰ Thời gian xuất: {now_str}  |  📊 Tổng sự kiện: {len(logs)}  |  🛠️ IT Tool LTT - ITTools (Lê Thế Tuấn - 0352 194 195)"
            sub_cell.font = Font(name="Segoe UI", size=9.5, italic=True, color="475569")
            sub_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            sub_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[2].height = 22

            # Row 3: Blank separator
            ws.row_dimensions[3].height = 8

            # 3. KPI Statistics Dashboard (Rows 4 & 5)
            ws.merge_cells("A4:L4")
            sec_stat = ws["A4"]
            sec_stat.value = "📌 BẢNG TỔNG HỢP TRẠNG THÁI & CHỈ SỐ HOẠT ĐỘNG NGUỒN ĐIỆN"
            sec_stat.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
            sec_stat.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
            sec_stat.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[4].height = 22

            kpi_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            # Row 5: KPI Summary Cards
            ws.row_dimensions[5].height = 26

            # Card 1: Tổng số
            ws.merge_cells("A5:B5")
            c_tot = ws["A5"]
            c_tot.value = f"Tổng sự kiện: {stats.get('total', len(logs))}"
            c_tot.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
            c_tot.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            c_tot.alignment = Alignment(horizontal="center", vertical="center")

            # Card 2: Tắt máy chủ động
            ws.merge_cells("C5:D5")
            c_shut = ws["C5"]
            c_shut.value = f"Tắt máy chủ động: {stats.get('shutdown_count', 0)}"
            c_shut.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")
            c_shut.fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
            c_shut.alignment = Alignment(horizontal="center", vertical="center")

            # Card 3: Khởi động lại
            ws.merge_cells("E5:F5")
            c_res = ws["E5"]
            c_res.value = f"Khởi động lại (Restart): {stats.get('restart_count', 0)}"
            c_res.font = Font(name="Segoe UI", size=10, bold=True, color="1D4ED8")
            c_res.fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
            c_res.alignment = Alignment(horizontal="center", vertical="center")

            # Card 4: Sập nguồn đột ngột (ID 6008)
            ws.merge_cells("G5:I5")
            c_unexp = ws["G5"]
            unexp_cnt = stats.get('unexpected_count', 0)
            c_unexp.value = f"Sập nguồn đột ngột (ID 6008): {unexp_cnt}"
            c_unexp.font = Font(name="Segoe UI", size=10, bold=True, color="92400E" if unexp_cnt > 0 else "475569")
            c_unexp.fill = PatternFill(start_color="FEF3C7" if unexp_cnt > 0 else "F8FAFC", end_color="FEF3C7" if unexp_cnt > 0 else "F8FAFC", fill_type="solid")
            c_unexp.alignment = Alignment(horizontal="center", vertical="center")

            # Card 5: Kernel-Power BSOD (ID 41)
            ws.merge_cells("J5:L5")
            c_kp = ws["J5"]
            kp_cnt = stats.get('kernel_power_count', 0)
            c_kp.value = f"BSOD / Treo sập (ID 41): {kp_cnt}"
            c_kp.font = Font(name="Segoe UI", size=10, bold=True, color="991B1B" if kp_cnt > 0 else "475569")
            c_kp.fill = PatternFill(start_color="FEE2E2" if kp_cnt > 0 else "F8FAFC", end_color="FEE2E2" if kp_cnt > 0 else "F8FAFC", fill_type="solid")
            c_kp.alignment = Alignment(horizontal="center", vertical="center")

            for col_idx in range(1, 13):
                ws.cell(row=5, column=col_idx).border = kpi_border

            # Row 6: Blank row
            ws.row_dimensions[6].height = 10

            # 4. Table Headers (Row 7)
            headers = [
                ("STT", 7, "center"),
                ("Thời Gian Ghi Nhận", 21, "center"),
                ("Event ID", 11, "center"),
                ("Mức Độ", 14, "center"),
                ("Loại Hành Động", 26, "left"),
                ("Người Dùng / Quyền", 24, "left"),
                ("Tiến Trình Phát Lệnh", 22, "left"),
                ("Lý Do Tắt Máy / Sự Cố", 30, "left"),
                ("Mã Reason Code", 16, "center"),
                ("Tên Máy Tính", 22, "left"),
                ("Chẩn Đoán Kỹ Thuật & Ghi Chú IT", 48, "left"),
                ("Chi Tiết Nhật Ký Gốc (Event Message)", 55, "left")
            ]

            ws.row_dimensions[7].height = 28
            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            for col_idx, (h_title, h_width, h_align) in enumerate(headers, 1):
                cell = ws.cell(row=7, column=col_idx, value=h_title)
                cell.font = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = h_width

            # 5. Populate Data Rows
            row_idx = 8
            for idx, item in enumerate(logs, start=1):
                bg_color = "FFFFFF" if idx % 2 != 0 else "F8FAFC"
                ws.row_dimensions[row_idx].height = 24

                eid = item.get("id")
                level = item.get("level") or ""
                action_lbl = item.get("action_label") or ""
                raw_msg = item.get("raw_message") or ""
                clean_raw = re.sub(r'[\r\n]+', '  |  ', raw_msg).strip()

                is_critical = (eid == 41 or 'Critical' in level)
                is_warning = (eid == 6008 or 'Warning' in level or 'Error' in level)
                is_restart = (item.get('action_type') == 'restart')

                row_values = [
                    (idx, "center", False, None, None),
                    (item.get("friendly_time") or item.get("time") or "", "center", False, None, None),
                    (eid, "center", True, "B91C1C" if is_critical else ("B45309" if is_warning else ("1D4ED8" if is_restart else "047857")), "FEE2E2" if is_critical else ("FEF3C7" if is_warning else ("EFF6FF" if is_restart else "ECFDF5"))),
                    (level, "center", True, "B91C1C" if is_critical else ("B45309" if is_warning else "334155"), "FEE2E2" if is_critical else ("FEF3C7" if is_warning else None)),
                    (action_lbl, "left", is_critical or is_warning, "B91C1C" if is_critical else ("B45309" if is_warning else "0F172A"), None),
                    (item.get("user") or "Hệ thống (SYSTEM)", "left", False, "334155", None),
                    (item.get("process_name") or item.get("process") or "-", "left", False, "334155", None),
                    (item.get("reason") or "Không có", "left", False, "0F172A", None),
                    (item.get("reason_code") or "-", "center", False, "64748B", None),
                    (item.get("machine") or comp_name, "left", False, "334155", None),
                    (item.get("expert_note") or "-", "left", False, "0F172A", None),
                    (clean_raw or "-", "left", False, "475569", None)
                ]

                for c_i, (val, align_h, is_bold, f_color, custom_bg) in enumerate(row_values, 1):
                    c = ws.cell(row=row_idx, column=c_i, value=val)
                    c_font_color = f_color if f_color else "0F172A"
                    c.font = Font(name="Segoe UI", size=9.5, bold=is_bold, color=c_font_color)
                    c_fill = custom_bg if custom_bg else bg_color
                    c.fill = PatternFill(start_color=c_fill, end_color=c_fill, fill_type="solid")
                    c.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=(c_i in (8, 11, 12)))
                    c.border = thin_border

                row_idx += 1

            # 6. Freeze Panes below headers (Row 8)
            ws.freeze_panes = "A8"

            # 7. Enable Auto-Filter on Table Headers (Row 7)
            if row_idx > 8:
                ws.auto_filter.ref = f"A7:L{row_idx - 1}"

            wb.save(save_path)
            return {
                "success": True,
                "message": f"✅ Đã xuất {len(logs)} bản ghi lịch sử ra file Excel (.xlsx) tuyệt đẹp tại:\n{save_path}",
                "file_path": save_path,
                "format": "excel"
            }
        except Exception as e:
            # Fallback to CSV if openpyxl fails
            pass

    # CSV Export format
    if not save_path or not save_path.lower().endswith('.csv'):
        save_path = os.path.join(desktop, f"Lich_Su_Tat_May_{ts}.csv")

    try:
        with open(save_path, mode='w', newline='', encoding='utf-8-sig') as f:
            # Tells Excel and WPS Office to explicitly divide columns using comma
            f.write("sep=,\n")
            writer = csv.writer(f)
            writer.writerow([
                "STT",
                "Thời Gian Ghi Nhận",
                "Event ID",
                "Mức Độ",
                "Loại Hành Động",
                "Người Dùng / Quyền",
                "Tiến Trình Phát Lệnh",
                "Lý Do Tắt Máy / Sự Cố",
                "Mã Reason Code",
                "Tên Máy Tính",
                "Chẩn Đoán Kỹ Thuật (Ghi Chú IT)",
                "Chi Tiết Nhật Ký Gốc"
            ])
            for idx, item in enumerate(logs, start=1):
                raw_msg = item.get("raw_message") or ""
                clean_raw = re.sub(r'[\r\n]+', ' | ', raw_msg).strip()
                clean_note = re.sub(r'[\r\n]+', ' | ', item.get("expert_note") or "").strip()
                clean_reason = re.sub(r'[\r\n]+', ' | ', item.get("reason") or "").strip()
                writer.writerow([
                    idx,
                    item.get("friendly_time") or item.get("time") or "",
                    item.get("id"),
                    item.get("level") or "",
                    item.get("action_label") or "",
                    item.get("user") or "Hệ thống (SYSTEM)",
                    item.get("process_name") or item.get("process") or "",
                    clean_reason,
                    item.get("reason_code") or "",
                    item.get("machine") or comp_name,
                    clean_note,
                    clean_raw
                ])
        return {
            "success": True,
            "message": f"✅ Đã xuất {len(logs)} bản ghi lịch sử ra file CSV có dấu tại:\n{save_path}",
            "file_path": save_path,
            "format": "csv"
        }
    except Exception as e:
        return {"success": False, "message": f"Lỗi khi lưu file CSV: {str(e)}"}



