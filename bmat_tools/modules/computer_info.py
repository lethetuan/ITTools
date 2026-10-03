"""
Computer Info Module
Shows RAM, CPU, GPU, HDD/SSD, Windows version, Battery, LAN info
"""

import tkinter as tk
from tkinter import ttk
import subprocess
import os
import sys
import threading
import platform
import ctypes

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from constants import COLORS, FONTS


def run_wmic(query):
    try:
        result = subprocess.run(
            ['wmic'] + query.split() + ['get', '/format:list'],
            capture_output=True, text=True, timeout=10
        )
        data = {}
        for line in result.stdout.splitlines():
            if '=' in line:
                k, _, v = line.partition('=')
                data[k.strip()] = v.strip()
        return data
    except:
        return {}


def run_ps(cmd):
    try:
        result = subprocess.run(
            ['powershell', '-Command', cmd],
            capture_output=True, text=True, timeout=15
        )
        return result.stdout.strip()
    except:
        return ''


class ComputerInfo:
    def __init__(self, parent):
        self.parent = parent
        self.parent.geometry('800x650')
        self.parent.configure(bg=COLORS['bg'])
        self.setup_ui()
        threading.Thread(target=self.load_info, daemon=True).start()

    def setup_ui(self):
        # Header
        hdr = tk.Frame(self.parent, bg=COLORS['bg_header'])
        hdr.pack(fill='x')
        tk.Label(hdr, text='💻  Computer Information', font=FONTS['large'],
                  bg=COLORS['bg_header'], fg='white', pady=10).pack(side='left', padx=15)
        tk.Button(hdr, text='🔄 Refresh', font=FONTS['small'],
                   bg=COLORS['accent'], fg='white', relief='flat',
                   padx=8, pady=3, cursor='hand2',
                   command=lambda: threading.Thread(
                       target=self.load_info, daemon=True).start()
                   ).pack(side='right', padx=15, pady=8)

        # Notebook
        self.nb = ttk.Notebook(self.parent)
        self.nb.pack(fill='both', expand=True, padx=8, pady=8)

        # Tabs
        self.tabs = {}
        for tab_name in ['📋 Summary', '🖥️ CPU', '💾 RAM', '🎮 GPU',
                          '💿 Storage', '🌐 Network', '🔋 Battery', '🪟 Windows']:
            f = tk.Frame(self.nb, bg=COLORS['bg'])
            self.nb.add(f, text=tab_name)
            self.tabs[tab_name] = f

        # Create content areas
        self.info_widgets = {}
        for tab_name, frame in self.tabs.items():
            text = tk.Text(frame, font=FONTS['mono'],
                            bg='#1E1E1E', fg='#D4D4D4',
                            insertbackground='white',
                            relief='flat', padx=15, pady=10,
                            state='disabled')
            vsb = ttk.Scrollbar(frame, command=text.yview)
            text.configure(yscrollcommand=vsb.set)
            vsb.pack(side='right', fill='y')
            text.pack(fill='both', expand=True)
            self.info_widgets[tab_name] = text

            # Configure tags
            text.tag_configure('header', foreground='#569CD6',
                                font=('Segoe UI', 11, 'bold'))
            text.tag_configure('key', foreground='#9CDCFE')
            text.tag_configure('value', foreground='#CE9178')
            text.tag_configure('good', foreground='#4EC9B0')
            text.tag_configure('warn', foreground='#DCDCAA')
            text.tag_configure('bad', foreground='#F44747')
            text.tag_configure('section', foreground='#C586C0',
                                font=('Consolas', 10, 'bold'))

        self.loading_label = tk.Label(self.parent,
                                       text='⏳ Loading system information...',
                                       font=FONTS['subtitle'],
                                       bg=COLORS['bg'], fg=COLORS['text_light'])
        self.loading_label.pack()

    def _write(self, widget, text, tag=None):
        widget.configure(state='normal')
        if tag:
            widget.insert('end', text, tag)
        else:
            widget.insert('end', text)
        widget.configure(state='disabled')

    def _clear(self, widget):
        widget.configure(state='normal')
        widget.delete('1.0', 'end')
        widget.configure(state='disabled')

    def _fmt_bytes(self, b):
        b = int(b)
        for unit in ['B', 'KB', 'MB', 'GB', 'TB']:
            if b < 1024:
                return f'{b:.1f} {unit}'
            b /= 1024
        return f'{b:.1f} PB'

    def load_info(self):
        self.loading_label.configure(text='⏳ Loading...')
        try:
            self._load_summary()
            self._load_cpu()
            self._load_ram()
            self._load_gpu()
            self._load_storage()
            self._load_network()
            self._load_battery()
            self._load_windows()
        except Exception as e:
            pass
        self.loading_label.configure(text='')

    def _load_summary(self):
        w = self.info_widgets['📋 Summary']
        self._clear(w)
        self._write(w, '═' * 60 + '\n', 'section')
        self._write(w, '          IT Tool LTT - System Summary\n', 'header')
        self._write(w, '═' * 60 + '\n\n', 'section')

        # Computer name
        comp_name = platform.node()
        username = os.environ.get('USERNAME', 'Unknown')
        self._write(w, f'  Computer Name : ', 'key')
        self._write(w, f'{comp_name}\n', 'value')
        self._write(w, f'  Username      : ', 'key')
        self._write(w, f'{username}\n', 'value')
        self._write(w, f'  Platform      : ', 'key')
        self._write(w, f'{platform.system()} {platform.release()}\n', 'value')
        self._write(w, f'  Architecture  : ', 'key')
        self._write(w, f'{platform.machine()}\n', 'value')

        # CPU
        self._write(w, '\n  ── CPU ──────────────────────────────────\n', 'section')
        cpu_data = run_wmic('cpu')
        self._write(w, f'  Model         : ', 'key')
        self._write(w, f'{cpu_data.get("Name", "Unknown")}\n', 'value')
        self._write(w, f'  Cores/Threads : ', 'key')
        cores = cpu_data.get('NumberOfCores', '?')
        threads = cpu_data.get('NumberOfLogicalProcessors', '?')
        self._write(w, f'{cores} Cores / {threads} Threads\n', 'value')

        # RAM
        self._write(w, '\n  ── RAM ──────────────────────────────────\n', 'section')
        mem_data = run_wmic('os')
        total_kb = int(mem_data.get('TotalVisibleMemorySize', 0))
        free_kb = int(mem_data.get('FreePhysicalMemory', 0))
        used_kb = total_kb - free_kb
        pct = (used_kb / total_kb * 100) if total_kb > 0 else 0
        self._write(w, f'  Total RAM     : ', 'key')
        self._write(w, f'{total_kb // 1024:,} MB ({total_kb // (1024*1024):,} GB)\n', 'value')
        self._write(w, f'  Used / Free   : ', 'key')
        tag = 'good' if pct < 70 else 'warn' if pct < 90 else 'bad'
        self._write(w, f'{used_kb // 1024:,} MB / {free_kb // 1024:,} MB ({pct:.0f}% used)\n', tag)

        # OS
        self._write(w, '\n  ── Windows ──────────────────────────────\n', 'section')
        winver = run_ps('(Get-WmiObject Win32_OperatingSystem).Caption')
        build = run_ps('(Get-WmiObject Win32_OperatingSystem).BuildNumber')
        self._write(w, f'  Version       : ', 'key')
        self._write(w, f'{winver}\n', 'value')
        self._write(w, f'  Build         : ', 'key')
        self._write(w, f'{build}\n', 'value')

    def _load_cpu(self):
        w = self.info_widgets['🖥️ CPU']
        self._clear(w)
        self._write(w, '  CPU Information\n\n', 'header')
        cpu = run_wmic('cpu')
        fields = [
            ('Name', 'Model'), ('Manufacturer', 'Manufacturer'),
            ('NumberOfCores', 'Physical Cores'),
            ('NumberOfLogicalProcessors', 'Logical Processors'),
            ('MaxClockSpeed', 'Max Clock Speed (MHz)'),
            ('CurrentClockSpeed', 'Current Clock Speed (MHz)'),
            ('L2CacheSize', 'L2 Cache (KB)'),
            ('L3CacheSize', 'L3 Cache (KB)'),
            ('Architecture', 'Architecture'),
            ('Caption', 'Caption'),
            ('LoadPercentage', 'CPU Load %'),
        ]
        for key, label in fields:
            val = cpu.get(key, 'N/A')
            if val:
                self._write(w, f'  {label:<30} : ', 'key')
                self._write(w, f'{val}\n', 'value')

        # Current usage via PowerShell
        usage = run_ps('(Get-WmiObject Win32_Processor).LoadPercentage')
        self._write(w, f'\n  {"Current Usage":<30} : ', 'key')
        pct = int(usage) if usage.isdigit() else 0
        tag = 'good' if pct < 60 else 'warn' if pct < 85 else 'bad'
        self._write(w, f'{pct}%\n', tag)

    def _load_ram(self):
        w = self.info_widgets['💾 RAM']
        self._clear(w)
        self._write(w, '  RAM Information\n\n', 'header')

        mem = run_wmic('os')
        total_kb = int(mem.get('TotalVisibleMemorySize', 0))
        free_kb = int(mem.get('FreePhysicalMemory', 0))
        used_kb = total_kb - free_kb

        self._write(w, f'  {"Total Physical RAM":<30} : ', 'key')
        self._write(w, f'{total_kb // 1024:,} MB\n', 'value')
        self._write(w, f'  {"Used":<30} : ', 'key')
        self._write(w, f'{used_kb // 1024:,} MB\n', 'warn')
        self._write(w, f'  {"Free":<30} : ', 'key')
        self._write(w, f'{free_kb // 1024:,} MB\n', 'good')

        # Physical slots
        self._write(w, '\n  Memory Slots:\n', 'section')
        result = subprocess.run(
            ['wmic', 'memorychip', 'get', 'BankLabel,Capacity,Speed,MemoryType,Manufacturer,PartNumber',
             '/format:csv'],
            capture_output=True, text=True
        )
        for line in result.stdout.splitlines()[2:]:
            parts = line.strip().split(',')
            if len(parts) >= 5 and parts[1]:
                try:
                    cap = int(parts[2]) // (1024 * 1024)
                    self._write(w, f'    Slot {parts[1]}: ', 'key')
                    self._write(w, f'{cap} MB  Speed: {parts[5]} MHz  {parts[3]}\n', 'value')
                except:
                    pass

    def _load_gpu(self):
        w = self.info_widgets['🎮 GPU']
        self._clear(w)
        self._write(w, '  GPU Information\n\n', 'header')
        gpu = run_wmic('path win32_VideoController')
        fields = [
            ('Name', 'GPU Model'),
            ('AdapterRAM', 'VRAM'),
            ('DriverVersion', 'Driver Version'),
            ('VideoProcessor', 'GPU Processor'),
            ('CurrentHorizontalResolution', 'Resolution Width'),
            ('CurrentVerticalResolution', 'Resolution Height'),
            ('CurrentRefreshRate', 'Refresh Rate (Hz)'),
            ('VideoModeDescription', 'Display Mode'),
        ]
        for key, label in fields:
            val = gpu.get(key, '')
            if val:
                if key == 'AdapterRAM':
                    try:
                        val = f'{int(val) // (1024*1024)} MB'
                    except:
                        pass
                self._write(w, f'  {label:<30} : ', 'key')
                self._write(w, f'{val}\n', 'value')

    def _load_storage(self):
        w = self.info_widgets['💿 Storage']
        self._clear(w)
        self._write(w, '  Storage Information\n\n', 'header')

        # Disk partitions
        result = subprocess.run(
            ['wmic', 'logicaldisk', 'get',
             'DeviceID,DriveType,Size,FreeSpace,FileSystem,VolumeName',
             '/format:csv'],
            capture_output=True, text=True
        )
        self._write(w, '  Drive Partitions:\n', 'section')
        for line in result.stdout.splitlines()[2:]:
            parts = line.strip().split(',')
            if len(parts) >= 6 and parts[2]:
                try:
                    drive = parts[1]
                    fs = parts[3]
                    free = int(parts[4]) if parts[4] else 0
                    size = int(parts[5]) if parts[5] else 0
                    vol = parts[6] if len(parts) > 6 else ''
                    if size > 0:
                        pct = (size - free) / size * 100
                        used = size - free
                        tag = 'good' if pct < 70 else 'warn' if pct < 90 else 'bad'
                        self._write(w, f'\n    Drive {drive}  [{vol}] ({fs})\n', 'key')
                        self._write(w,
                                    f'      Total: {self._fmt_bytes(size)}'
                                    f'  Used: {self._fmt_bytes(used)}'
                                    f'  Free: {self._fmt_bytes(free)}'
                                    f'  ({pct:.0f}% full)\n', tag)
                except:
                    pass

        # Physical disks
        self._write(w, '\n  Physical Disks:\n', 'section')
        disk_info = subprocess.run(
            ['wmic', 'diskdrive', 'get', 'Caption,Size,MediaType,InterfaceType',
             '/format:csv'],
            capture_output=True, text=True
        )
        for line in disk_info.stdout.splitlines()[2:]:
            parts = line.strip().split(',')
            if len(parts) >= 5 and parts[2]:
                try:
                    size = int(parts[4]) if parts[4] else 0
                    self._write(w, f'    {parts[2]}\n', 'key')
                    self._write(w,
                                f'      Size: {self._fmt_bytes(size)}'
                                f'  Type: {parts[3]}  Interface: {parts[1]}\n', 'value')
                except:
                    pass

    def _load_network(self):
        w = self.info_widgets['🌐 Network']
        self._clear(w)
        self._write(w, '  Network Information\n\n', 'header')

        # Get all adapters
        result = run_ps(
            'Get-NetIPConfiguration | ForEach-Object {'
            '"Adapter: " + $_.InterfaceAlias + "`n"'
            ' + "  IPv4: " + ($_.IPv4Address.IPAddress -join ", ") + "`n"'
            ' + "  IPv6: " + ($_.IPv6Address.IPAddress -join ", ") + "`n"'
            ' + "  Gateway: " + ($_.IPv4DefaultGateway.NextHop -join ", ") + "`n"'
            ' + "  DNS: " + ($_.DNSServer.ServerAddresses -join ", ") + "`n"'
            '}'
        )
        if result:
            self._write(w, result + '\n', 'value')

        # MAC addresses
        self._write(w, '\n  MAC Addresses:\n', 'section')
        mac_result = run_ps(
            'Get-NetAdapter | Select-Object Name, MacAddress, Status, LinkSpeed '
            '| Format-Table -AutoSize | Out-String'
        )
        self._write(w, mac_result + '\n', 'value')

        # External IP
        self._write(w, '\n  External IP:\n', 'section')
        try:
            import urllib.request
            ext_ip = urllib.request.urlopen('https://api.ipify.org', timeout=5).read().decode()
            self._write(w, f'  {ext_ip}\n', 'good')
        except:
            self._write(w, '  Cannot retrieve (no internet)\n', 'bad')

    def _load_battery(self):
        w = self.info_widgets['🔋 Battery']
        self._clear(w)
        self._write(w, '  Battery Information\n\n', 'header')

        bat = run_wmic('path Win32_Battery')
        if not bat:
            self._write(w, '  No battery found (Desktop PC)\n', 'warn')
            return

        fields = [
            ('Name', 'Battery Name'),
            ('Status', 'Status'),
            ('BatteryStatus', 'Charge Status'),
            ('EstimatedChargeRemaining', 'Charge Remaining (%)'),
            ('EstimatedRunTime', 'Estimated Runtime (min)'),
            ('DesignCapacity', 'Design Capacity (mWh)'),
            ('FullChargeCapacity', 'Full Charge Capacity (mWh)'),
        ]
        for key, label in fields:
            val = bat.get(key, '')
            if val:
                self._write(w, f'  {label:<35} : ', 'key')
                if key == 'EstimatedChargeRemaining':
                    pct = int(val) if val.isdigit() else 0
                    tag = 'good' if pct > 40 else 'warn' if pct > 15 else 'bad'
                    self._write(w, f'{val}%\n', tag)
                else:
                    self._write(w, f'{val}\n', 'value')

        # Power plan
        self._write(w, '\n  Power Plan:\n', 'section')
        plan = run_ps('powercfg /getactivescheme')
        self._write(w, f'  {plan}\n', 'value')

    def _load_windows(self):
        w = self.info_widgets['🪟 Windows']
        self._clear(w)
        self._write(w, '  Windows Information\n\n', 'header')

        fields_ps = [
            ('(Get-WmiObject Win32_OperatingSystem).Caption', 'OS Name'),
            ('(Get-WmiObject Win32_OperatingSystem).Version', 'Version'),
            ('(Get-WmiObject Win32_OperatingSystem).BuildNumber', 'Build Number'),
            ('(Get-WmiObject Win32_OperatingSystem).OSArchitecture', 'Architecture'),
            ('(Get-WmiObject Win32_OperatingSystem).InstallDate', 'Install Date'),
            ('(Get-WmiObject Win32_OperatingSystem).LastBootUpTime', 'Last Boot'),
            ('(Get-WmiObject Win32_OperatingSystem).RegisteredUser', 'Registered User'),
            ('(Get-WmiObject Win32_OperatingSystem).SerialNumber', 'Serial Number'),
        ]
        for cmd, label in fields_ps:
            val = run_ps(cmd)
            if val:
                self._write(w, f'  {label:<30} : ', 'key')
                self._write(w, f'{val}\n', 'value')

        # Activation status
        self._write(w, '\n  Activation:\n', 'section')
        act = run_ps('(Get-WmiObject SoftwareLicensingProduct | Where-Object { $_.PartialProductKey }).LicenseStatus')
        status_map = {'1': 'Licensed ✅', '0': 'Unlicensed ❌'}
        self._write(w, f'  {status_map.get(act, act)}\n', 'value')

        # Product key (partial)
        self._write(w, '\n  Product Key (partial):\n', 'section')
        key = run_ps('(Get-WmiObject SoftwareLicensingProduct | Where-Object { $_.PartialProductKey }).PartialProductKey')
        self._write(w, f'  xxxxx-xxxxx-xxxxx-xxxxx-{key}\n', 'value')


def get_system_power_status():
    """Returns accurate real-time power status using Win32 GetSystemPowerStatus."""
    import sys
    if sys.platform != 'win32':
        return None
    try:
        import ctypes
        from ctypes import wintypes
        class SYSTEM_POWER_STATUS(ctypes.Structure):
            _fields_ = [
                ('ACLineStatus', wintypes.BYTE),
                ('BatteryFlag', wintypes.BYTE),
                ('BatteryLifePercent', wintypes.BYTE),
                ('SystemStatusFlag', wintypes.BYTE),
                ('BatteryLifeTime', wintypes.DWORD),
                ('BatteryFullLifeTime', wintypes.DWORD),
            ]
        s = SYSTEM_POWER_STATUS()
        if ctypes.windll.kernel32.GetSystemPowerStatus(ctypes.byref(s)):
            ac_status = s.ACLineStatus # 0: Offline, 1: Online, 255: Unknown
            bat_flag = s.BatteryFlag   # 128: No battery, 8: Charging
            bat_pct = s.BatteryLifePercent # 0..100, 255: Unknown
            has_battery = (not bool(bat_flag & 128)) and (bat_pct != 255)
            is_charging = bool(bat_flag & 8)
            is_ac = (ac_status == 1)

            if not has_battery:
                status_text = "Không có pin (Desktop PC)"
            elif is_charging:
                status_text = "Đang sạc (AC Powered)"
            elif is_ac:
                if bat_pct >= 99:
                    status_text = "Đầy 100% (Đang cắm sạc)"
                else:
                    status_text = "Đang cắm sạc (AC Powered)"
            else:
                status_text = "Đang dùng Pin (Battery Powered)"

            return {
                "has_battery": has_battery,
                "is_laptop": has_battery,
                "is_ac": is_ac,
                "is_charging": is_charging,
                "level_pct": bat_pct if bat_pct <= 100 else 100,
                "status_text": status_text
            }
    except Exception:
        pass
    return None


def get_device_manager_issues():
    """
    Scans for genuine missing, failed, or warning drivers using Windows PnP API.
    Filters out phantom (disconnected USB/Bluetooth) and user-disabled devices.
    Returns accurate real-time list matching Windows Device Manager.
    """
    import subprocess
    import json

    error_code_map = {
        1: "Chưa được cấu hình đúng (Code 1)",
        3: "Driver bị hỏng hoặc thiếu tài nguyên (Code 3)",
        10: "Thiết bị không thể khởi động (Code 10)",
        12: "Xung đột tài nguyên phần cứng (Code 12)",
        14: "Cần khởi động lại máy để hoàn tất cài đặt (Code 14)",
        18: "Cần cài đặt lại driver cho thiết bị (Code 18)",
        19: "Cấu hình Registry của thiết bị bị lỗi (Code 19)",
        24: "Thiết bị không hiện diện hoặc hoạt động không đúng (Code 24)",
        28: "Chưa cài đặt Driver (Code 28 - Missing Driver)",
        29: "Thiết bị bị vô hiệu hóa trong BIOS/Firmware (Code 29)",
        31: "Windows không thể tải các driver cho thiết bị (Code 31)",
        32: "Dịch vụ của driver đã bị tắt (Code 32)",
        37: "Windows không thể khởi tạo driver thiết bị (Code 37)",
        38: "Phiên bản driver trước vẫn còn trong bộ nhớ (Code 38)",
        39: "Driver bị hỏng hoặc thiếu file cài đặt (Code 39)",
        43: "Windows đã dừng thiết bị do phát hiện sự cố (Code 43)",
        52: "Driver chưa được ký chữ ký số hợp lệ (Code 52)"
    }

    issues = []
    try:
        # Query only physically present devices that have real error states:
        # Excludes:
        # - ConfigManagerErrorCode == 0 (OK)
        # - ConfigManagerErrorCode == 22 (Manually disabled by user, not an error)
        # - ConfigManagerErrorCode == 45 / Problem == 'CM_PROB_PHANTOM' (Unplugged/disconnected USB/Bluetooth devices)
        cmd = (
            "Get-PnpDevice -PresentOnly | "
            "Where-Object { "
            "($_.ConfigManagerErrorCode -ne 0 -and $_.ConfigManagerErrorCode -ne 22 -and $_.ConfigManagerErrorCode -ne 45) -or "
            "($_.Problem -ne 'CM_PROB_NONE' -and $_.Problem -ne 'CM_PROB_PHANTOM' -and $_.Problem -ne 'CM_PROB_DISABLED') -or "
            "($_.Status -eq 'Error') "
            "} | Select-Object FriendlyName, Class, InstanceId, Problem, ConfigManagerErrorCode, Status"
        )
        res = subprocess.run(
            ['powershell', '-NoProfile', '-Command', f"{cmd} | ConvertTo-Json -Depth 2"],
            capture_output=True, text=True, encoding='utf-8', errors='ignore', timeout=10
        )
        if res.stdout.strip():
            parsed = json.loads(res.stdout.strip())
            if isinstance(parsed, dict):
                parsed = [parsed]
            for dev in parsed:
                fname = (dev.get("FriendlyName") or dev.get("Name") or "").strip()
                if not fname:
                    continue
                err_code = dev.get("ConfigManagerErrorCode", 0)
                try:
                    err_code = int(err_code)
                except (ValueError, TypeError):
                    err_code = 0

                desc = error_code_map.get(err_code)
                if not desc:
                    prob = dev.get("Problem") or ""
                    if prob and prob not in ("CM_PROB_NONE", "CM_PROB_PHANTOM", "CM_PROB_DISABLED"):
                        desc = f"Lỗi thiết bị ({prob})"
                    elif err_code:
                        desc = f"Lỗi thiết bị (Code {err_code})"
                    else:
                        desc = "Cảnh báo / Lỗi Driver"

                issues.append({
                    "name": fname,
                    "class": (dev.get("Class") or "Unknown").strip(),
                    "instance": (dev.get("InstanceId") or dev.get("DeviceID") or "").strip(),
                    "error_code": err_code,
                    "error_desc": desc,
                    "status": dev.get("Status", "Error")
                })
    except Exception:
        pass

    return issues


def get_detailed_hardware_info():
    """Gathers comprehensive hardware & system info matching BTP Tool Pro 2026 inspection dashboard."""
    import subprocess
    import json
    import socket
    import datetime
    import os

    def fmt_bytes_gb(b):
        try:
            val = float(b)
            return f"{val / (1024**3):.1f} GB"
        except Exception:
            return "N/A"

    def run_ps_cmd(cmd):
        try:
            res = subprocess.run(
                ['powershell', '-NoProfile', '-Command', f"{cmd} | ConvertTo-Json -Depth 3"],
                capture_output=True, text=True, encoding='utf-8', errors='ignore'
            )
            if res.stdout.strip():
                parsed = json.loads(res.stdout.strip())
                if isinstance(parsed, dict):
                    return [parsed]
                elif isinstance(parsed, list):
                    return parsed
            return []
        except Exception:
            return []

    # 1. CPU
    cpu_list = run_ps_cmd('Get-CimInstance Win32_Processor | Select-Object Name, NumberOfCores, NumberOfLogicalProcessors, MaxClockSpeed, SocketDesignation, L2CacheSize, L3CacheSize')
    cpu_info = {
        "name": "N/A",
        "cores_threads": "N/A",
        "max_clock": "N/A",
        "socket": "N/A",
        "cache": "N/A"
    }
    if cpu_list:
        c = cpu_list[0]
        cpu_info["name"] = c.get("Name", "N/A").strip()
        cores = c.get("NumberOfCores", 0)
        threads = c.get("NumberOfLogicalProcessors", 0)
        cpu_info["cores_threads"] = f"{cores} nhân · {threads} luồng"
        clock_mhz = c.get("MaxClockSpeed", 0)
        cpu_info["max_clock"] = f"{clock_mhz / 1000.0:.2f} GHz" if clock_mhz else "N/A"
        cpu_info["socket"] = c.get("SocketDesignation", "N/A").strip() or "N/A"
        l2 = c.get("L2CacheSize", 0)
        l3 = c.get("L3CacheSize", 0)
        cache_strs = []
        if l2: cache_strs.append(f"{l2/1024:.1f} MB L2")
        if l3: cache_strs.append(f"{l3/1024:.1f} MB L3")
        cpu_info["cache"] = " · ".join(cache_strs) if cache_strs else "N/A"

    # 2. System & BaseBoard & BIOS
    cs_list = run_ps_cmd('Get-CimInstance Win32_ComputerSystem | Select-Object Manufacturer, Model, Name, TotalPhysicalMemory, SystemFamily, ChassisSKUNumber')
    board_list = run_ps_cmd('Get-CimInstance Win32_BaseBoard | Select-Object Manufacturer, Product, SerialNumber, Version')
    bios_list = run_ps_cmd('Get-CimInstance Win32_BIOS | Select-Object Manufacturer, SMBIOSBIOSVersion, ReleaseDate, SerialNumber')
    prod_list = run_ps_cmd('Get-CimInstance Win32_ComputerSystemProduct | Select-Object IdentifyingNumber, UUID, Name, Vendor')

    system_info = {
        "vendor": "N/A",
        "manufacturer": "N/A",
        "model": "N/A",
        "computer_name": socket.gethostname(),
        "system_family": "N/A",
        "chassis_type": "Desktop",
        "is_laptop": False
    }
    ram_total_str = "N/A"
    ram_total_gb_num = 0.0

    if cs_list:
        cs = cs_list[0]
        mfg = cs.get("Manufacturer", "").strip()
        mdl = cs.get("Model", "").strip()
        system_info["manufacturer"] = mfg or "N/A"
        system_info["model"] = mdl or "N/A"
        system_info["vendor"] = f"{mfg} {mdl}".strip() or "N/A"
        system_info["computer_name"] = cs.get("Name", socket.gethostname()).strip()
        system_info["system_family"] = cs.get("SystemFamily", "").strip() or mdl or "N/A"
        if cs.get("TotalPhysicalMemory"):
            ram_b = float(cs.get("TotalPhysicalMemory", 0))
            ram_total_gb_num = round(ram_b / (1024**3), 1)
            ram_total_str = f"{ram_total_gb_num} GB"

    # Service Tag & UUID
    service_tag_info = {
        "serial": "N/A",
        "service_tag": "N/A",
        "product_id": "N/A",
        "uuid": "N/A"
    }

    if prod_list:
        p = prod_list[0]
        ident = p.get("IdentifyingNumber", "").strip()
        uuid_val = p.get("UUID", "").strip()
        if ident and ident.lower() != "default string":
            service_tag_info["serial"] = ident
            service_tag_info["service_tag"] = ident
            service_tag_info["product_id"] = ident
        if uuid_val:
            service_tag_info["uuid"] = uuid_val

    if bios_list:
        bi = bios_list[0]
        bios_serial = bi.get("SerialNumber", "").strip()
        if bios_serial and bios_serial.lower() != "default string":
            if service_tag_info["serial"] == "N/A":
                service_tag_info["serial"] = bios_serial
                service_tag_info["service_tag"] = bios_serial
                service_tag_info["product_id"] = bios_serial

    # Mainboard
    mainboard_info = {
        "manufacturer": "N/A",
        "model": "N/A",
        "serial": "N/A",
        "version": "N/A"
    }

    if board_list:
        b = board_list[0]
        mainboard_info["manufacturer"] = b.get("Manufacturer", "N/A").strip()
        mainboard_info["model"] = b.get("Product", "N/A").strip()
        mainboard_info["serial"] = b.get("SerialNumber", "N/A").strip() or "N/A"
        mainboard_info["version"] = b.get("Version", "A00").strip() or "A00"

    # BIOS Info
    bios_info = {
        "vendor": "N/A",
        "version": "N/A",
        "release_date": "N/A"
    }

    if bios_list:
        bi = bios_list[0]
        bios_info["vendor"] = bi.get("Manufacturer", "N/A").strip()
        bios_info["version"] = bi.get("SMBIOSBIOSVersion", "N/A").strip()
        b_date = bi.get("ReleaseDate", "")
        date_clean = ""
        if b_date and isinstance(b_date, str) and "Date(" in b_date:
            try:
                ms = int(b_date.split("(")[1].split(")")[0])
                date_clean = datetime.datetime.fromtimestamp(ms / 1000.0).strftime("%Y-%m-%d")
            except Exception:
                date_clean = ""
        elif b_date and isinstance(b_date, str) and len(b_date) >= 8:
            date_clean = f"{b_date[:4]}-{b_date[4:6]}-{b_date[6:8]}"
        bios_info["release_date"] = date_clean or "N/A"

    # 3. Battery Details & Health (real-time Win32 power status + CIM metadata)
    p_status = get_system_power_status()
    bat_list = run_ps_cmd('Get-CimInstance Win32_Battery | Select-Object Name, DeviceID, EstimatedChargeRemaining, BatteryStatus, DesignCapacity, FullChargeCapacity')
    battery_info = {
        "is_laptop": False,
        "name": "N/A",
        "level_pct": 100,
        "wear_pct": 0.0,
        "design_mwh": "N/A",
        "full_mwh": "N/A",
        "status_text": "Không có pin (Desktop PC)",
        "health_text": "PC Desktop"
    }

    if p_status and p_status["has_battery"]:
        battery_info["is_laptop"] = True
        system_info["is_laptop"] = True
        system_info["chassis_type"] = "Notebook"
        battery_info["level_pct"] = p_status["level_pct"]
        battery_info["status_text"] = p_status["status_text"]

    if bat_list:
        bt = bat_list[0]
        battery_info["is_laptop"] = True
        system_info["is_laptop"] = True
        system_info["chassis_type"] = "Notebook"
        battery_info["name"] = bt.get("Name", bt.get("DeviceID", "Standard Battery")).strip()
        if not (p_status and p_status["has_battery"]):
            battery_info["level_pct"] = bt.get("EstimatedChargeRemaining", 100)
            bat_status = bt.get("BatteryStatus", 1)
            if bat_status in [2, 6, 7, 8, 9]:
                battery_info["status_text"] = "Đang sạc (AC Powered)"
            elif bat_status == 3:
                battery_info["status_text"] = "Đầy 100% (Đang cắm sạc)"
            else:
                battery_info["status_text"] = "Đang dùng Pin (Battery Powered)"

        design_cap = float(bt.get("DesignCapacity", 0) or 0)
        full_cap = float(bt.get("FullChargeCapacity", 0) or 0)

        if design_cap > 0 and full_cap > 0:
            battery_info["design_mwh"] = f"{int(design_cap)} mWh"
            battery_info["full_mwh"] = f"{int(full_cap)} mWh"
            wear = round(max(0.0, 100.0 - (full_cap / design_cap * 100.0)), 1)
            battery_info["wear_pct"] = wear
            if wear < 20:
                battery_info["health_text"] = "✓ Pin tốt"
            elif wear < 40:
                battery_info["health_text"] = "⚠️ Chai nhẹ"
            else:
                battery_info["health_text"] = "❌ Chai nặng (Nên thay)"
        else:
            battery_info["health_text"] = "✓ Hoạt động tốt"

    # 4. RAM Modules
    ram_list = run_ps_cmd('Get-CimInstance Win32_PhysicalMemory | Select-Object DeviceLocator, Capacity, Speed, Manufacturer, PartNumber')
    ram_modules = []
    for r in ram_list:
        loc = r.get("DeviceLocator", "DIMM").strip()
        cap = fmt_bytes_gb(r.get("Capacity", 0))
        speed = f"{r.get('Speed', 0)} MHz" if r.get('Speed') else ""
        mfg = r.get("Manufacturer", "").strip()
        part = r.get("PartNumber", "").strip()
        info_str = " · ".join(filter(None, [cap, speed, mfg, part]))
        ram_modules.append({
            "locator": loc,
            "details": info_str
        })

    # 5. GPU
    gpu_list = run_ps_cmd('Get-CimInstance Win32_VideoController | Select-Object Name, AdapterRAM, DriverVersion')
    gpus = []
    for idx, g in enumerate(gpu_list, 1):
        name = g.get("Name", "").strip()
        if not name:
            continue
        vram_b = g.get("AdapterRAM", 0)
        vram_str = fmt_bytes_gb(vram_b) if vram_b else ""
        driver = f"driver {g.get('DriverVersion', '').strip()}" if g.get('DriverVersion') else ""
        info_str = " · ".join(filter(None, [name, vram_str, driver]))
        gpus.append({
            "label": f"GPU {idx}",
            "name": name,
            "vram": vram_str,
            "driver": g.get("DriverVersion", "").strip(),
            "details": info_str
        })

    # 6. Physical Disks & Partitions
    disk_list = run_ps_cmd('Get-CimInstance Win32_DiskDrive | Select-Object Model, Size, MediaType')
    disks = []
    for idx, d in enumerate(disk_list, 1):
        model = d.get("Model", "").strip()
        if not model:
            continue
        size_str = fmt_bytes_gb(d.get("Size", 0))
        media = d.get("MediaType", "SSD / Fixed Disk").strip()
        info_str = " · ".join(filter(None, [model, size_str, media]))
        disks.append({
            "label": f"Ổ {idx}",
            "model": model,
            "size": size_str,
            "details": info_str
        })

    # Partitions (C:, D:, E:...)
    logical_disks = run_ps_cmd('Get-CimInstance Win32_LogicalDisk -Filter "DriveType=3" | Select-Object DeviceID, FreeSpace, Size, VolumeName')
    partitions = []
    for ld in logical_disks:
        dev = ld.get("DeviceID", "C:").strip()
        free_b = float(ld.get("FreeSpace", 0))
        total_b = float(ld.get("Size", 0))
        free_gb = round(free_b / (1024**3), 1)
        total_gb = round(total_b / (1024**3), 1)
        partitions.append({
            "drive": dev,
            "free_gb": free_gb,
            "total_gb": total_gb,
            "label": f"{dev} {free_gb} GB tự do / {total_gb} GB"
        })

    # 7. OS Info
    os_list = run_ps_cmd('Get-CimInstance Win32_OperatingSystem | Select-Object Caption, Version, BuildNumber, OSArchitecture, InstallDate')
    os_info = {
        "caption": "Windows 11 Pro 64-bit",
        "version": "",
        "build": "",
        "arch": "64-bit",
        "install_date": "N/A"
    }
    if os_list:
        o = os_list[0]
        os_info["caption"] = o.get("Caption", "Windows 11 Pro").strip()
        os_info["version"] = o.get("Version", "").strip()
        os_info["build"] = o.get("BuildNumber", "").strip()
        os_info["arch"] = o.get("OSArchitecture", "64-bit").strip()
        inst_date = o.get("InstallDate", "")
        if inst_date and isinstance(inst_date, str) and len(inst_date) >= 8:
            os_info["install_date"] = f"{inst_date[:4]}-{inst_date[4:6]}-{inst_date[6:8]}"

    # 8. Missing / Warning Drivers Check (Real-time & accurate, matching Device Manager)
    missing_drivers = get_device_manager_issues()

    return {
        "success": True,
        "cpu": cpu_info,
        "system": system_info,
        "service_tag": service_tag_info,
        "mainboard": mainboard_info,
        "bios_info": bios_info,
        "battery": battery_info,
        "ram_total": ram_total_str,
        "ram_modules": ram_modules,
        "gpus": gpus,
        "disks": disks,
        "partitions": partitions,
        "os": os_info,
        "missing_drivers": missing_drivers
    }


def export_specs(format_type="xlsx", specs_data=None, output_dir=None):
    """
    Exports full, comprehensive computer specs to Excel (.xlsx) or CSV (.csv).
    Includes all hardware, CPU, RAM modules, GPUs, Disks, Partitions, Battery, OS, Network, etc.
    """
    import datetime
    import csv
    import os
    import subprocess

    if not specs_data:
        specs_data = get_detailed_hardware_info()

    user_profile = os.environ.get('USERPROFILE', 'C:\\')
    if not output_dir or not os.path.exists(output_dir):
        output_dir = os.path.join(user_profile, 'Desktop')

    sys_info = specs_data.get("system", {})
    comp_name = sys_info.get("computer_name", "PC")
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

    # Collect network info if possible
    net_adapters = []
    try:
        import modules.ip_manager as im
        net_adapters = im.get_network_adapters().get("adapters", [])
    except Exception:
        pass

    # Build full list of rows: (Category, Attribute, Value)
    rows = []

    # 1. System Info
    sec = "1. THÔNG TIN MÁY TÍNH & HỆ THỐNG"
    rows.append((sec, "Tên Máy Tính (Computer Name)", comp_name))
    rows.append((sec, "Hãng Sản Xuất (Manufacturer)", sys_info.get("manufacturer", "N/A")))
    rows.append((sec, "Model Máy (System Model)", sys_info.get("model", "N/A")))
    rows.append((sec, "Dòng Sản Phẩm (System Family)", sys_info.get("system_family", "N/A")))
    rows.append((sec, "Kiểu Dáng (Chassis Type)", "Laptop / Notebook" if sys_info.get("is_laptop") else "Desktop PC"))
    rows.append((sec, "Service Tag / Serial Number", specs_data.get("service_tag", {}).get("service_tag", "N/A")))
    rows.append((sec, "Mã Định Danh Duy Nhất (UUID)", specs_data.get("service_tag", {}).get("uuid", "N/A")))

    # 2. CPU
    sec = "2. BỘ XỬ LÝ (CPU)"
    cpu = specs_data.get("cpu", {})
    rows.append((sec, "Tên Vi Xử Lý (CPU Name)", cpu.get("name", "N/A")))
    rows.append((sec, "Số Nhân & Số Luồng", cpu.get("cores_threads", "N/A")))
    rows.append((sec, "Xung Nhịp Tối Đa (Max Clock)", cpu.get("max_clock", "N/A")))
    rows.append((sec, "Socket CPU", cpu.get("socket", "N/A")))
    rows.append((sec, "Bộ Nhớ Đệm (Cache)", cpu.get("cache", "N/A")))

    # 3. Mainboard & BIOS
    sec = "3. BO MẠCH CHỦ (MAINBOARD) & BIOS"
    mb = specs_data.get("mainboard", {})
    bios = specs_data.get("bios_info", {})
    rows.append((sec, "Hãng Sản Xuất Mainboard", mb.get("manufacturer", "N/A")))
    rows.append((sec, "Model Mainboard", mb.get("model", "N/A")))
    rows.append((sec, "Serial Mainboard", mb.get("serial", "N/A")))
    rows.append((sec, "Phiên Bản Mainboard (Revision)", mb.get("version", "N/A")))
    rows.append((sec, "Nhà Cung Cấp BIOS (Vendor)", bios.get("vendor", "N/A")))
    rows.append((sec, "Phiên Bản BIOS (Version)", bios.get("version", "N/A")))
    rows.append((sec, "Ngày Phát Hành BIOS (Release Date)", bios.get("release_date", "N/A")))

    # 4. RAM
    sec = "4. BỘ NHỚ RAM (MEMORY)"
    rows.append((sec, "Tổng Dung Lượng RAM", specs_data.get("ram_total", "N/A")))
    ram_modules = specs_data.get("ram_modules", [])
    if ram_modules:
        for idx, r in enumerate(ram_modules, 1):
            rows.append((sec, f"Khe Cắm RAM #{idx} ({r.get('locator', 'Slot')})", r.get("details", "N/A")))
    else:
        rows.append((sec, "Chi Tiết Khe Cắm", "Không lấy được thông tin khe cắm"))

    # 5. GPU
    sec = "5. CARD ĐỒ HỌA (GPU)"
    gpus = specs_data.get("gpus", [])
    if gpus:
        for idx, g in enumerate(gpus, 1):
            rows.append((sec, f"Card Đồ Họa #{idx} ({g.get('label', 'GPU')})", g.get("details", g.get("name", "N/A"))))
    else:
        rows.append((sec, "Card Đồ Họa", "N/A"))

    # 6. Physical Disks
    sec = "6. Ổ ĐĨA VẬT LÝ (PHYSICAL DISKS)"
    disks = specs_data.get("disks", [])
    if disks:
        for idx, d in enumerate(disks, 1):
            rows.append((sec, f"Ổ Cứng Vật Lý #{idx} ({d.get('label', 'Ổ')})", d.get("details", d.get("model", "N/A"))))
    else:
        rows.append((sec, "Ổ Cứng", "N/A"))

    # 7. Partitions
    sec = "7. PHÂN VÙNG Ổ ĐĨA (PARTITIONS)"
    partitions = specs_data.get("partitions", [])
    if partitions:
        for p in partitions:
            total = p.get("total_gb", 0)
            free = p.get("free_gb", 0)
            used = round(total - free, 1) if total >= free else 0
            pct = round((used / total * 100), 1) if total > 0 else 0
            rows.append((sec, f"Phân Vùng {p.get('drive', '')}", f"Tổng: {total} GB | Đã dùng: {used} GB ({pct}%) | Còn trống: {free} GB"))
    else:
        rows.append((sec, "Phân Vùng", "N/A"))

    # 8. Battery
    sec = "8. THÔNG TIN PIN & NGUỒN ĐIỆN"
    bat = specs_data.get("battery", {})
    if bat.get("is_laptop"):
        rows.append((sec, "Tên Pin (Device Name)", bat.get("name", "Standard Battery")))
        rows.append((sec, "Mức Pin Hiện Tại", f"{bat.get('level_pct', 0)}%"))
        rows.append((sec, "Độ Chai Pin (Wear Level)", f"{bat.get('wear_pct', 0)}% ({bat.get('health_text', '')})"))
        rows.append((sec, "Dung Lượng Thiết Kế", bat.get("design_mwh", "N/A")))
        rows.append((sec, "Dung Lượng Sạc Đầy Thực Tế", bat.get("full_mwh", "N/A")))
        rows.append((sec, "Trạng Thái Nguồn Điện", bat.get("status_text", "N/A")))
    else:
        rows.append((sec, "Tình Trạng Pin", "Không có pin (Máy tính bàn - Desktop PC)"))

    # 9. Operating System
    sec = "9. HỆ ĐIỀU HÀNH (WINDOWS)"
    os_info = specs_data.get("os", {})
    rows.append((sec, "Tên Hệ Điều Hành", os_info.get("caption", "Windows")))
    rows.append((sec, "Phiên Bản (Version)", os_info.get("version", "N/A")))
    rows.append((sec, "Số Bản Build (Build Number)", str(os_info.get("build", "N/A"))))
    rows.append((sec, "Kiến Trúc Hệ Thống", os_info.get("arch", "64-bit")))
    rows.append((sec, "Ngày Cài Đặt Windows", os_info.get("install_date", "N/A")))

    # 10. Network
    sec = "10. KẾT NỐI MẠNG (NETWORK)"
    if net_adapters:
        for idx, a in enumerate(net_adapters, 1):
            ip_str = a.get("ip", "N/A")
            mac_str = a.get("mac", "N/A")
            gw_str = a.get("gateway", "N/A")
            dns_str = " / ".join(a.get("dns", [])) if a.get("dns") else "N/A"
            desc = f"IP: {ip_str} | MAC: {mac_str} | Gateway: {gw_str} | DNS: {dns_str}"
            rows.append((sec, f"Card Mạng #{idx}: {a.get('name', 'Adapter')}", desc))
    else:
        rows.append((sec, "Card Mạng", "N/A"))

    # 11. Missing Drivers
    sec = "11. DRIVER CẢNH BÁO / THIẾU"
    missing = specs_data.get("missing_drivers", [])
    if missing:
        for idx, m in enumerate(missing, 1):
            rows.append((sec, f"Thiết Bị Cảnh Báo #{idx}", f"{m.get('name')} | Lớp: {m.get('class')} | {m.get('error_desc', 'Lỗi Driver')}"))
    else:
        rows.append((sec, "Tình Trạng Driver", "✓ Hoạt động hoàn hảo - 100% thiết bị có driver đầy đủ"))

    fmt = format_type.lower()
    if fmt == "csv":
        file_name = f"CauHinh_{comp_name}_{ts}.csv"
        file_path = os.path.join(output_dir, file_name)
        with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["STT", "Danh Mục", "Thuộc Tính / Thiết Bị", "Thông Số Chi Tiết"])
            for idx, (cat, prop, val) in enumerate(rows, 1):
                writer.writerow([idx, cat, prop, val])
        return {"success": True, "file_path": file_path, "message": f"✅ Đã xuất toàn bộ cấu hình ra file CSV thành công tại:\n{file_path}"}

    else:
        # Excel .xlsx format
        file_name = f"CauHinh_{comp_name}_{ts}.xlsx"
        file_path = os.path.join(output_dir, file_name)
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Cấu Hình Chi Tiết"

            # Title Banner
            ws.merge_cells("A1:D1")
            title_cell = ws["A1"]
            title_cell.value = f"BÁO CÁO CẤU HÌNH HỆ THỐNG MÁY TÍNH - {comp_name.upper()}"
            title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
            title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 36

            # Subtitle Banner
            ws.merge_cells("A2:D2")
            sub_cell = ws["A2"]
            sub_cell.value = f"Thời gian xuất: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M:%S')}  |  Phần mềm: IT Tool LTT 2026 - Tác giả: Lê Thế Tuấn (0352 194 195)"
            sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="475569")
            sub_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            sub_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[2].height = 22

            # Table Headers (Row 4)
            headers = ["STT", "Danh Mục", "Thuộc Tính / Thiết Bị", "Thông Số Chi Tiết"]
            ws.row_dimensions[4].height = 26
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=4, column=col_idx)
                cell.value = h
                cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
                cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "left", vertical="center")

            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            current_cat = None
            row_idx = 5
            item_num = 1

            for cat, prop, val in rows:
                if cat != current_cat:
                    current_cat = cat
                    # Section Header Row
                    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
                    sec_cell = ws.cell(row=row_idx, column=1)
                    sec_cell.value = cat
                    sec_cell.font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
                    sec_cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
                    sec_cell.alignment = Alignment(horizontal="left", vertical="center")
                    ws.row_dimensions[row_idx].height = 24
                    row_idx += 1

                bg_color = "FFFFFF" if item_num % 2 != 0 else "F8FAFC"
                ws.row_dimensions[row_idx].height = 20

                # Col A: STT
                cA = ws.cell(row=row_idx, column=1, value=item_num)
                cA.alignment = Alignment(horizontal="center", vertical="center")
                cA.font = Font(name="Segoe UI", size=10, color="64748B")

                # Col B: Danh Mục
                cB = ws.cell(row=row_idx, column=2, value=cat.split(".", 1)[-1].strip())
                cB.alignment = Alignment(horizontal="left", vertical="center")
                cB.font = Font(name="Segoe UI", size=10, color="334155")

                # Col C: Thuộc Tính
                cC = ws.cell(row=row_idx, column=3, value=prop)
                cC.alignment = Alignment(horizontal="left", vertical="center")
                cC.font = Font(name="Segoe UI", size=10, bold=True, color="1E293B")

                # Col D: Giá Trị
                cD = ws.cell(row=row_idx, column=4, value=val)
                cD.alignment = Alignment(horizontal="left", vertical="center")
                cD.font = Font(name="Segoe UI", size=10, color="0F172A")

                for c in [cA, cB, cC, cD]:
                    c.border = thin_border
                    c.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

                row_idx += 1
                item_num += 1

            # Auto-adjust column widths
            ws.column_dimensions['A'].width = 8
            ws.column_dimensions['B'].width = 34
            ws.column_dimensions['C'].width = 38
            ws.column_dimensions['D'].width = 70

            # Freeze panes below headers
            ws.freeze_panes = "A5"

            wb.save(file_path)
            return {"success": True, "file_path": file_path, "message": f"✅ Đã xuất toàn bộ cấu hình ra file Excel (.xlsx) thành công tại:\n{file_path}"}
        except Exception as e:
            # Fallback to CSV if openpyxl fails for any reason
            csv_name = f"CauHinh_{comp_name}_{ts}.csv"
            csv_path = os.path.join(output_dir, csv_name)
            with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(["STT", "Danh Mục", "Thuộc Tính / Thiết Bị", "Thông Số Chi Tiết"])
                for idx, (cat, prop, val) in enumerate(rows, 1):
                    writer.writerow([idx, cat, prop, val])
            return {"success": True, "file_path": csv_path, "message": f"✅ Đã xuất cấu hình ra file CSV tại:\n{csv_path}"}

